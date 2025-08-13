#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import re
from pathlib import Path
from typing import List, Tuple, Optional

import numpy as np
import pandas as pd
import statsmodels.api as sm
import matplotlib.pyplot as plt
import seaborn as sns
from docx import Document

sns.set(style="whitegrid", font="DejaVu Sans")

# 工作目录与输出路径
WORKDIR = Path("/workspace")
OUTPUT_DIR = WORKDIR / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
FONTS_DIR = WORKDIR / "fonts"
FONTS_DIR.mkdir(parents=True, exist_ok=True)
CHS_FONT_PATH = FONTS_DIR / "NotoSansSC-Regular.otf"

# 文件路径
FILE_CITY_EMISSIONS = WORKDIR / "1997-2019年290个中国城市碳排放清单.xlsx"
FILE_APPARENT = WORKDIR / "表观碳排放清单_1997-2021.xlsx"
FILE_NEI_GBA = WORKDIR / "粤港澳大湾区新能源产业规模 的副本.xls"

# 九市名单（标准化后的城市名）
GBA_CITIES_CN = ["广州", "深圳", "珠海", "佛山", "惠州", "东莞", "中山", "江门", "肇庆"]


# ------------------------ 工具函数 ------------------------

def strip_zwsp(text: str) -> str:
    if text is None:
        return text
    return re.sub("\u200b", "", str(text)).strip()


def ensure_chinese_font() -> Optional[str]:
    try:
        from matplotlib import font_manager, rcParams
        # 优先使用项目内字体
        local_font = CHS_FONT_PATH
        if not local_font.exists():
            import urllib.request
            url = (
                "https://github.com/googlefonts/noto-cjk/raw/main/Sans/OTF/SimplifiedChinese/"
                "NotoSansSC-Regular.otf"
            )
            try:
                urllib.request.urlretrieve(url, local_font)
            except Exception:
                pass
        if local_font.exists():
            font_manager.fontManager.addfont(str(local_font))
        # 常见系统中文字体候选
        candidates = [
            "Noto Sans CJK SC",  # Noto 家族内部名
            "Noto Sans SC",
            "Microsoft YaHei",  # Windows雅黑
            "SimHei",            # Windows黑体
            "PingFang SC",       # macOS苹方
            "Source Han Sans CN",# 思源黑体
            "WenQuanYi Zen Hei", # Linux文泉驿
            "DejaVu Sans",       # 最后兜底
        ]
        rcParams["font.family"] = ["sans-serif"]
        rcParams["font.sans-serif"] = candidates
        rcParams["axes.unicode_minus"] = False
        # 刷新字体缓存，确保新字体生效
        try:
            font_manager._load_fontmanager(try_read_cache=False)  # type: ignore[attr-defined]
        except Exception:
            pass
        return str(local_font) if local_font.exists() else None
    except Exception:
        return None


def normalize_city_cn(name: str) -> str:
    s = strip_zwsp(name)
    # 去省名前缀
    for prov in [
        "广东", "广西", "江苏", "浙江", "山东", "河南", "河北", "湖南", "湖北", "江西", "福建", "云南", "贵州", "四川",
        "安徽", "山西", "辽宁", "吉林", "黑龙江", "内蒙古", "青海", "甘肃", "陕西", "宁夏", "新疆", "海南",
        "北京", "天津", "上海", "重庆",
    ]:
        if s.startswith(prov):
            s = s[len(prov) :]
            break
    # 去城市后缀
    for suf in ["市", "地区", "盟", "自治州", "特别行政区", "区"]:
        if s.endswith(suf):
            s = s[: -len(suf)]
    return s


def to_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series.astype(str).str.replace(r"[^0-9eE+\-.]", "", regex=True), errors="coerce")


def safe_log(series: pd.Series, eps: float = 1e-6) -> pd.Series:
    s = pd.to_numeric(series, errors="coerce")
    return np.log(s.clip(lower=eps))


# ------------------------ 读数函数 ------------------------

def load_city_emissions() -> pd.DataFrame:
    xls = pd.ExcelFile(FILE_CITY_EMISSIONS, engine="openpyxl")
    df = pd.read_excel(xls, sheet_name="emission vector")
    df.columns = [c.strip() for c in df.columns]
    df = df[["city", "year", "emission"]]
    df["city"] = df["city"].astype(str).str.strip()
    df["city_std"] = df["city"].apply(normalize_city_cn)
    df = df[df["city_std"].isin(GBA_CITIES_CN)].copy()
    df["year"] = pd.to_numeric(df["year"], errors="coerce").astype(int)
    df["emission"] = pd.to_numeric(df["emission"], errors="coerce")
    return df


def aggregate_gba_emissions_city_to_year(df_city: pd.DataFrame) -> pd.DataFrame:
    agg = (
        df_city.groupby("year", as_index=False)["emission"].sum().rename(columns={"emission": "emission_gba"})
    )
    return agg.sort_values("year")


def load_apparent_emissions_guangdong() -> pd.DataFrame:
    xls = pd.ExcelFile(FILE_APPARENT, engine="openpyxl")
    years = [s for s in xls.sheet_names if re.fullmatch(r"\d{4}", str(s))]
    records = []
    for y in years:
        dfy = pd.read_excel(xls, sheet_name=y)
        dfy.columns = [strip_zwsp(c) for c in dfy.columns]
        col_gd = None
        for cand in ["Guangdong", "广东", "Guang Dong"]:
            if cand in dfy.columns:
                col_gd = cand
                break
        if not col_gd or "Items" not in dfy.columns:
            continue
        mask = dfy["Items"].isna() & dfy[col_gd].notna()
        if mask.any():
            val = float(pd.to_numeric(dfy.loc[mask, col_gd]).astype(float).iloc[0])
            records.append({"year": int(y), "emission_guangdong": val})
    return pd.DataFrame.from_records(records).sort_values("year")


def load_gba_nei() -> pd.DataFrame:
    df = pd.read_excel(FILE_NEI_GBA, sheet_name=0)
    df.columns = [strip_zwsp(c) for c in df.columns]
    rename_map = {
        "年份": "year",
        "可再生能源装机总量（GW）": "re_total_gw",
        "光伏装机（GW）": "pv_gw",
        "风电装机（GW）": "wind_gw",
        "新能源汽车保有量（万辆）": "nev_10k_units",
        "储能装机（MW）": "storage_mw",
        "关键事件": "events",
    }
    # 容错匹配
    for key in list(rename_map.keys()):
        if key not in df.columns:
            for col in df.columns:
                if key.replace("（", "(").replace("）", ")") in col.replace("（", "(").replace("）", ")"):
                    rename_map[col] = rename_map.pop(key)
                    break
    df = df.rename(columns=rename_map)
    for col in ["year", "re_total_gw", "pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
        if col in df.columns:
            df[col] = to_numeric(df[col])
    df = df[~df["year"].isna()].copy()
    df["year"] = df["year"].astype(int)
    return df.sort_values("year")


# ------------------------ 建模辅助 ------------------------

def fit_time_series_ols(
    df: pd.DataFrame, y: str, x_vars: List[str], trend: bool = True, cluster: Optional[pd.Series] = None
) -> sm.regression.linear_model.RegressionResultsWrapper:
    X = df[x_vars].copy()
    if trend:
        X["trend"] = np.arange(len(df)) + 1
    X = sm.add_constant(X)
    model = sm.OLS(df[y].astype(float), X.astype(float), missing="drop")
    if cluster is not None:
        res = model.fit(cov_type="cluster", cov_kwds={"groups": cluster})
    else:
        res = model.fit(cov_type="HAC", cov_kwds={"maxlags": 1})
    return res


def add_city_dummies(df: pd.DataFrame, city_col: str) -> Tuple[pd.DataFrame, List[str]]:
    dummies = pd.get_dummies(df[city_col], prefix="city", drop_first=True)
    return pd.concat([df, dummies], axis=1), list(dummies.columns)


# ------------------------ 图表与导出 ------------------------

def compute_descriptive_stats(df: pd.DataFrame, cols: List[str]) -> pd.DataFrame:
    sub = df[cols].copy()
    desc = sub.describe(percentiles=[0.25, 0.5, 0.75]).T
    desc = desc.rename(columns={
        "count": "样本数",
        "mean": "均值",
        "std": "标准差",
        "min": "最小值",
        "25%": "P25",
        "50%": "中位数",
        "75%": "P75",
        "max": "最大值",
    })
    desc.insert(0, "变量", desc.index)
    return desc.reset_index(drop=True)


def save_corr_heatmap(df: pd.DataFrame, cols: List[str], title: str, out_png: Path) -> None:
    ensure_chinese_font()
    corr = df[cols].corr()
    plt.figure(figsize=(6, 5))
    sns.heatmap(corr, annot=True, fmt=".2f", cmap="Blues", square=True)
    plt.title(title)
    plt.tight_layout()
    plt.savefig(out_png, dpi=200)
    plt.close()


def plot_and_export(df_gba_year: pd.DataFrame, df_nei: pd.DataFrame, merged_gba: pd.DataFrame, panel: pd.DataFrame, outputs):
    ensure_chinese_font()

    # 图1：九市碳排放总量
    fig1, ax1 = plt.subplots(figsize=(10, 5))
    ax1.plot(df_gba_year["year"], df_gba_year["emission_gba"], marker="o", color="#2E86DE")
    ax1.axvspan(2014, 2019, color="#2E86DE", alpha=0.08, label="与新能源指标重叠区间")
    ax1.set_title("粤港澳大湾区九市碳排放总量（1997-2019）")
    ax1.set_xlabel("年份")
    ax1.set_ylabel("碳排放（原单位）")
    ax1.legend()
    fig1.tight_layout()
    fig1.savefig(OUTPUT_DIR / "图1_九市碳排放总量_1997_2019.png", dpi=200)
    plt.close(fig1)

    # 图2：新能源装机规模
    fig2, ax2 = plt.subplots(figsize=(10, 6))
    ax2.plot(df_nei["year"], df_nei["pv_gw"], marker="o", label="光伏装机（GW）")
    ax2.plot(df_nei["year"], df_nei["wind_gw"], marker="o", label="风电装机（GW）")
    ax2.plot(df_nei["year"], df_nei["re_total_gw"], marker="o", label="可再生能源装机总量（GW）")
    ax2.set_title("大湾区新能源装机规模（2014-2023）")
    ax2.set_xlabel("年份")
    ax2.set_ylabel("装机规模（GW）")
    ax2.legend()
    fig2.tight_layout()
    fig2.savefig(OUTPUT_DIR / "图2_新能源装机规模_2014_2023.png", dpi=200)
    plt.close(fig2)

    # 图3：新能源汽车与储能
    fig3, ax3a = plt.subplots(figsize=(10, 6))
    ax3a.plot(df_nei["year"], df_nei["nev_10k_units"], color="#27AE60", marker="s", label="新能源汽车保有量（万辆）")
    ax3a.set_xlabel("年份")
    ax3a.set_ylabel("新能源汽车保有量（万辆）", color="#27AE60")
    ax3a.tick_params(axis="y", labelcolor="#27AE60")
    ax3b = ax3a.twinx()
    ax3b.plot(df_nei["year"], df_nei["storage_mw"], color="#8E44AD", marker="^", label="储能装机（MW）")
    ax3b.set_ylabel("储能装机（MW）", color="#8E44AD")
    ax3b.tick_params(axis="y", labelcolor="#8E44AD")
    ax3a.set_title("大湾区新能源汽车与储能发展（2014-2023）")
    fig3.tight_layout()
    fig3.savefig(OUTPUT_DIR / "图3_NEVE与储能_2014_2023.png", dpi=200)
    plt.close(fig3)

    # 图4：散点关系
    if not merged_gba.empty:
        labels_map = {
            "pv_gw": "光伏装机（GW）",
            "wind_gw": "风电装机（GW）",
            "nev_10k_units": "新能源汽车保有量（万辆）",
            "storage_mw": "储能装机（MW）",
        }
        for var, label in labels_map.items():
            fig, ax = plt.subplots(figsize=(6, 5))
            sns.regplot(data=merged_gba, x=var, y="emission_gba", ax=ax, marker="o", color="#2C3E50")
            ax.set_title(f"九市碳排放与{label}的关系（2014-2019）")
            ax.set_xlabel(label)
            ax.set_ylabel("九市碳排放（原单位）")
            fig.tight_layout()
            fig.savefig(OUTPUT_DIR / f"图4_散点_{var}_vs_九市碳排放_2014_2019.png", dpi=200)
            plt.close(fig)

    # 图5：双轴图
    if not merged_gba.empty:
        fig4, ax4a = plt.subplots(figsize=(10, 6))
        ax4a.plot(merged_gba["year"], merged_gba["emission_gba"], color="#e67e22", marker="o", label="九市碳排放")
        ax4a.set_xlabel("年份")
        ax4a.set_ylabel("九市碳排放（原单位）", color="#e67e22")
        ax4a.tick_params(axis="y", labelcolor="#e67e22")
        ax4b = ax4a.twinx()
        ax4b.plot(merged_gba["year"], merged_gba["re_total_gw"], color="#2980b9", marker="s", label="可再生能源装机总量（GW）")
        ax4b.set_ylabel("可再生能源装机总量（GW）", color="#2980b9")
        ax4b.tick_params(axis="y", labelcolor="#2980b9")
        ax4a.set_title("九市碳排放与可再生能源装机（2014-2019）")
        fig4.tight_layout()
        fig4.savefig(OUTPUT_DIR / "图5_双轴_九市碳排放_vs_可再生装机_2014_2019.png", dpi=200)
        plt.close(fig4)

    # 图6：城市趋势
    if not panel.empty:
        fig5, ax5 = plt.subplots(figsize=(10, 6))
        for c, sub in panel.groupby("city_std"):
            ax5.plot(sub["year"], sub["emission"], marker="o", label=c)
        ax5.set_title("九市城市碳排放趋势（2014-2019）")
        ax5.set_xlabel("年份")
        ax5.set_ylabel("碳排放（原单位）")
        ax5.legend(ncols=3, fontsize=8)
        fig5.tight_layout()
        fig5.savefig(OUTPUT_DIR / "图6_九市城市碳排放趋势_2014_2019.png", dpi=200)
        plt.close(fig5)

    # 导出回归结果（CSV与Word）
    def tidy_result(res: sm.regression.linear_model.RegressionResultsWrapper, name: str) -> pd.DataFrame:
        params = res.params
        se = res.bse
        conf = res.conf_int()
        pvals = res.pvalues
        out = pd.DataFrame(
            {
                "变量": params.index,
                "系数": params.values,
                "标准误": se.values,
                "下界95%": conf[0].values,
                "上界95%": conf[1].values,
                "p值": pvals.values,
            }
        )
        out.insert(0, "模型", name)
        return out

    all_tidy = []
    for name, res in outputs:
        all_tidy.append(tidy_result(res, name))
    if all_tidy:
        tidy_df = pd.concat(all_tidy, ignore_index=True)
        tidy_df.to_csv(OUTPUT_DIR / "回归结果_可粘贴.csv", index=False)
        doc = Document()
        doc.add_heading("回归结果汇总", level=1)
        for name, res in outputs:
            doc.add_heading(name, level=2)
            tbl = doc.add_table(rows=1, cols=6)
            hdr = tbl.rows[0].cells
            hdr[0].text = "变量"; hdr[1].text = "系数"; hdr[2].text = "标准误"; hdr[3].text = "下界95%"; hdr[4].text = "上界95%"; hdr[5].text = "p值"
            params = res.params
            se = res.bse
            conf = res.conf_int()
            p = res.pvalues
            for var in params.index:
                row = tbl.add_row().cells
                row[0].text = str(var)
                row[1].text = f"{params[var]:.4f}"
                row[2].text = f"{se[var]:.4f}"
                row[3].text = f"{conf.loc[var, 0]:.4f}"
                row[4].text = f"{conf.loc[var, 1]:.4f}"
                row[5].text = f"{p[var]:.4f}"
        doc.save(OUTPUT_DIR / "回归结果汇总.docx")


def plot_coef_forest(res: sm.regression.linear_model.RegressionResultsWrapper, name: str, out_path: Path, title: Optional[str] = None):
    ensure_chinese_font()
    params = res.params.copy()
    conf = res.conf_int()
    # 仅绘制核心解释变量（对数项）
    mask = [v for v in params.index if v.startswith("ln_")]
    if not mask:
        return
    coef = params[mask]
    ci_low = conf.loc[mask, 0]
    ci_high = conf.loc[mask, 1]
    order = coef.abs().sort_values(ascending=True).index
    plt.figure(figsize=(7, max(2, 0.5*len(order)+1)))
    y = np.arange(len(order))
    plt.hlines(y, ci_low[order], ci_high[order], color="#2C3E50")
    plt.plot(coef[order], y, 'o', color="#e67e22")
    plt.axvline(0, color="#7f8c8d", lw=1)
    plt.yticks(y, order)
    plt.xlabel("系数及95%置信区间")
    plt.title(title or f"{name}：回归系数图")
    plt.tight_layout()
    plt.savefig(out_path, dpi=200)
    plt.close()


def plot_corr_table_bars(csv_path: Path, label_prefix: str):
    ensure_chinese_font()
    if not csv_path.exists():
        return
    import pandas as pd
    df = pd.read_csv(csv_path)
    # 针对不同类型绘制Pearson条形图
    vars_ = df['变量'].unique()
    types = df['类型'].unique()
    for v in vars_:
        sub = df[df['变量']==v]
        plt.figure(figsize=(6,4))
        plt.bar(sub['类型'], sub['Pearson'], color="#3498db")
        plt.axhline(0, color="#7f8c8d", lw=1)
        plt.ylabel('Pearson相关系数')
        plt.xticks(rotation=20)
        plt.title(f"{label_prefix}：{v} 与排放的相关（水平/差分/去趋势）")
        plt.tight_layout()
        plt.savefig(OUTPUT_DIR / f"图_相关性条形_{label_prefix}_{v}.png", dpi=200)
        plt.close()


def elasticity_table_from_outputs(outputs: List[Tuple[str, sm.regression.linear_model.RegressionResultsWrapper]]) -> pd.DataFrame:
    rows = []
    for name, res in outputs:
        for var in res.params.index:
            if var.startswith("ln_"):
                beta = float(res.params[var])
                rows.append({
                    "模型": name,
                    "变量": var,
                    "弹性(系数)": beta,
                    "10%变动对应排放变化(%)": beta * 10.0,
                    "5%变动对应排放变化(%)": beta * 5.0,
                })
    return pd.DataFrame(rows)


def first_difference(df: pd.DataFrame, y: str, x_vars: List[str], group_col: Optional[str] = None) -> pd.DataFrame:
    d = df.copy().sort_values([group_col, "year"]) if group_col else df.copy().sort_values("year")
    for col in [y] + x_vars:
        d[f"d_{col}"] = d.groupby(group_col)[col].diff() if group_col else d[col].diff()
    if group_col:
        keep_cols = [group_col, "year"] + [f"d_{y}"] + [f"d_{x}" for x in x_vars]
    else:
        keep_cols = ["year"] + [f"d_{y}"] + [f"d_{x}" for x in x_vars]
    return d[keep_cols].dropna().reset_index(drop=True)


def add_year_dummies(df: pd.DataFrame) -> Tuple[pd.DataFrame, List[str]]:
    year_d = pd.get_dummies(df["year"].astype(int), prefix="year", drop_first=True)
    return pd.concat([df, year_d], axis=1), list(year_d.columns)


def robustness_checks(merged_gba: pd.DataFrame, merged_gd: pd.DataFrame, panel_merged: pd.DataFrame) -> List[Tuple[str, sm.regression.linear_model.RegressionResultsWrapper]]:
    outs: List[Tuple[str, sm.regression.linear_model.RegressionResultsWrapper]] = []

    # -- 九市时间序列：无趋势、差分、滞后 --
    if len(merged_gba) >= 5:
        d = merged_gba.copy().sort_values("year")
        d["ln_emission_gba"] = safe_log(d["emission_gba"])
        for v in ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
            d[f"ln_{v}"] = safe_log(d[v])
        xvars = ["ln_pv_gw", "ln_wind_gw", "ln_nev_10k_units", "ln_storage_mw"]
        # 无趋势
        outs.append(("九市-无趋势（2014-2019）", fit_time_series_ols(d, "ln_emission_gba", xvars, trend=False)))
        # 差分
        fd = first_difference(d, "ln_emission_gba", xvars)
        if len(fd) >= 4:
            outs.append(("九市-一阶差分（2014-2019）", fit_time_series_ols(fd.rename(columns={"d_ln_emission_gba": "y"}), "y", [f"d_{x}" for x in xvars], trend=False)))
        # 滞后X
        d_lag = d.copy()
        for v in xvars:
            d_lag[v + "_lag1"] = d_lag[v].shift(1)
        d_lag = d_lag.dropna().reset_index(drop=True)
        if len(d_lag) >= 5:
            outs.append(("九市-滞后解释变量（t-1）", fit_time_series_ols(d_lag, "ln_emission_gba", [v + "_lag1" for v in xvars], trend=True)))

    # -- 广东时间序列：无趋势、差分、滞后 --
    if len(merged_gd) >= 6:
        d = merged_gd.copy().sort_values("year")
        d["ln_emission_gd"] = safe_log(d["emission_guangdong"])
        for v in ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
            d[f"ln_{v}"] = safe_log(d[v])
        xvars = ["ln_pv_gw", "ln_wind_gw", "ln_nev_10k_units", "ln_storage_mw"]
        outs.append(("广东-无趋势（2014-2021）", fit_time_series_ols(d, "ln_emission_gd", xvars, trend=False)))
        fd = first_difference(d, "ln_emission_gd", xvars)
        if len(fd) >= 5:
            outs.append(("广东-一阶差分（2014-2021）", fit_time_series_ols(fd.rename(columns={"d_ln_emission_gd": "y"}), "y", [f"d_{x}" for x in xvars], trend=False)))
        d_lag = d.copy()
        for v in xvars:
            d_lag[v + "_lag1"] = d_lag[v].shift(1)
        d_lag = d_lag.dropna().reset_index(drop=True)
        if len(d_lag) >= 6:
            outs.append(("广东-滞后解释变量（t-1）", fit_time_series_ols(d_lag, "ln_emission_gd", [v + "_lag1" for v in xvars], trend=True)))

    # -- 面板：年虚拟变量替代趋势 --
    if len(panel_merged) >= 30:
        d = panel_merged.copy()
        d["ln_emission_city"] = safe_log(d["emission"])
        for v in ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
            d[f"ln_{v}"] = safe_log(d[v])
        d, city_dummies = add_city_dummies(d, "city_std")
        d, year_dummies = add_year_dummies(d)
        xvars = ["ln_pv_gw", "ln_wind_gw", "ln_nev_10k_units", "ln_storage_mw"] + city_dummies + year_dummies
        outs.append(("九市-城市FE+年份FE（2014-2019）", fit_time_series_ols(d, "ln_emission_city", xvars, trend=False, cluster=d["year"])) )

    return outs


def correlation_table_levels_diffs_detrend(df: pd.DataFrame, y_col: str, x_cols: List[str], label_prefix: str) -> pd.DataFrame:
    """构造水平、对数差分、去趋势残差的相关性表（Pearson/Spearman）。"""
    rows = []
    d = df.copy().sort_values("year")
    # 使用对数
    d[f"ln_{y_col}"] = safe_log(d[y_col])
    for x in x_cols:
        d[f"ln_{x}"] = safe_log(d[x])
    # 去趋势残差
    t = np.arange(len(d)) + 1
    for var in [f"ln_{y_col}"] + [f"ln_{x}" for x in x_cols]:
        X = sm.add_constant(pd.Series(t, name="trend"))
        res = sm.OLS(d[var], X).fit()
        d[var+"_res"] = d[var] - res.fittedvalues
    # 一阶差分
    d_diff = d[["year"] + [f"ln_{y_col}"] + [f"ln_{x}" for x in x_cols]].diff().dropna()

    def add_row(kind: str, xname: str, s1: pd.Series, s2: pd.Series):
        rows.append({
            "口径": label_prefix,
            "类型": kind,
            "变量": xname,
            "Pearson": float(s1.corr(s2, method="pearson")),
            "Spearman": float(s1.corr(s2, method="spearman")),
            "样本数": int(s1.dropna().shape[0])
        })

    # 水平（对数）
    for x in x_cols:
        add_row("水平(对数)", x, d[f"ln_{y_col}"], d[f"ln_{x}"])
    # 差分（对数差分）
    for x in x_cols:
        add_row("一阶差分(对数)", x, d_diff[f"ln_{y_col}"], d_diff[f"ln_{x}"])
    # 去趋势（残差）
    for x in x_cols:
        add_row("去趋势(残差)", x, d[f"ln_{y_col}_res"], d[f"ln_{x}_res"])

    return pd.DataFrame(rows)


def lead_lag_cross_corr(df: pd.DataFrame, y_col: str, x_cols: List[str], label_prefix: str, lags: List[int] = [-2, -1, 0, 1, 2]) -> pd.DataFrame:
    d = df.copy().sort_values("year")
    d[f"ln_{y_col}"] = safe_log(d[y_col])
    for x in x_cols:
        d[f"ln_{x}"] = safe_log(d[x])
    rows = []
    for x in x_cols:
        for L in lags:
            if L >= 0:
                s_x = d[f"ln_{x}"] .shift(L)
                s_y = d[f"ln_{y_col}"]
            else:
                # 负滞后：y领先
                s_x = d[f"ln_{x}"]
                s_y = d[f"ln_{y_col}"] .shift(-L)
            s = pd.concat([s_x, s_y], axis=1).dropna()
            if len(s) >= 3:
                corr = float(s.iloc[:,0].corr(s.iloc[:,1]))
            else:
                corr = np.nan
            rows.append({"口径": label_prefix, "变量": x, "滞后L": L, "相关系数": corr, "样本数": len(s)})
    return pd.DataFrame(rows)


def plot_lead_lag(df_ccf: pd.DataFrame, label_prefix: str, out_prefix: str):
    ensure_chinese_font()
    for var, sub in df_ccf.groupby("变量"):
        plt.figure(figsize=(6,4))
        plt.axhline(0, color="#999", lw=1)
        markerline, stemlines, baseline = plt.stem(sub["滞后L"], sub["相关系数"], linefmt="C0-", markerfmt="C0o", basefmt="C7-")
        try:
            plt.setp(stemlines, linewidth=1.5)
        except Exception:
            pass
        plt.title(f"{label_prefix}：{var} 与排放的领先-滞后相关")
        plt.xlabel("滞后 L（正值表示X领先排放L年）")
        plt.ylabel("相关系数")
        plt.tight_layout()
        plt.savefig(OUTPUT_DIR / f"{out_prefix}_leadlag_{var}.png", dpi=200)
        plt.close()


def pairplot_levels(df: pd.DataFrame, cols_map: dict, title: str, out_png: Path):
    ensure_chinese_font()
    sub = df[list(cols_map.keys())].rename(columns=cols_map).copy()
    g = sns.pairplot(sub, diag_kind="kde")
    plt.suptitle(title, y=1.02)
    plt.tight_layout()
    g.savefig(out_png, dpi=200)
    plt.close()


def residualized_scatter(df: pd.DataFrame, y_col: str, x_col: str, label_prefix: str, out_png: Path):
    ensure_chinese_font()
    d = df.copy().sort_values("year")
    # 对数并去趋势
    for col in [y_col, x_col]:
        d[f"ln_{col}"] = safe_log(d[col])
        X = sm.add_constant(pd.Series(np.arange(len(d))+1, name="trend"))
        res = sm.OLS(d[f"ln_{col}"], X).fit()
        d[col+"_res"] = d[f"ln_{col}"] - res.fittedvalues
    plt.figure(figsize=(5,4))
    sns.regplot(x=d[x_col+"_res"], y=d[y_col+"_res"], marker="o")
    plt.title(f"{label_prefix}：去趋势后残差散点（{x_col} vs {y_col}）")
    plt.xlabel(f"{x_col}（去趋势残差）")
    plt.ylabel(f"{y_col}（去趋势残差）")
    plt.tight_layout()
    plt.savefig(out_png, dpi=200)
    plt.close()


def preliminary_association_analysis(merged_gba: pd.DataFrame, merged_gd: pd.DataFrame):
    """两者关联性初步分析：
    - 水平/差分/去趋势相关性（Pearson/Spearman）
    - 领先-滞后相关性（-2..+2）与折线图
    - 水平散点矩阵（pairplot）
    - 去趋势残差散点（示例：光伏/风电）
    """
    x_cols = ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]
    # 九市
    if not merged_gba.empty:
        corr_tbl = correlation_table_levels_diffs_detrend(merged_gba, "emission_gba", x_cols, "九市")
        corr_tbl.to_csv(OUTPUT_DIR / "关联性_相关性表_九市.csv", index=False)
        # 领先滞后
        ccf = lead_lag_cross_corr(merged_gba, "emission_gba", x_cols, "九市")
        ccf.to_csv(OUTPUT_DIR / "关联性_领先滞后相关_九市.csv", index=False)
        plot_lead_lag(ccf, "九市", "九市")
        # Pairplot（水平）
        cols_map = {
            "emission_gba": "九市碳排放",
            "pv_gw": "光伏装机(GW)",
            "wind_gw": "风电装机(GW)",
            "nev_10k_units": "NEV(万辆)",
            "storage_mw": "储能(MW)",
        }
        pairplot_levels(merged_gba, cols_map, "九市：新能源与碳排放的散点矩阵", OUTPUT_DIR / "图_散点矩阵_九市.png")
        # 去趋势残差散点（示例）
        residualized_scatter(merged_gba, "emission_gba", "pv_gw", "九市", OUTPUT_DIR / "图_去趋势残差散点_九市_pv.png")
        residualized_scatter(merged_gba, "emission_gba", "wind_gw", "九市", OUTPUT_DIR / "图_去趋势残差散点_九市_wind.png")

    # 广东
    if not merged_gd.empty:
        corr_tbl = correlation_table_levels_diffs_detrend(merged_gd, "emission_guangdong", x_cols, "广东")
        corr_tbl.to_csv(OUTPUT_DIR / "关联性_相关性表_广东.csv", index=False)
        ccf = lead_lag_cross_corr(merged_gd, "emission_guangdong", x_cols, "广东")
        ccf.to_csv(OUTPUT_DIR / "关联性_领先滞后相关_广东.csv", index=False)
        plot_lead_lag(ccf, "广东", "广东")
        cols_map = {
            "emission_guangdong": "广东碳排放",
            "pv_gw": "光伏装机(GW)",
            "wind_gw": "风电装机(GW)",
            "nev_10k_units": "NEV(万辆)",
            "storage_mw": "储能(MW)",
        }
        pairplot_levels(merged_gd, cols_map, "广东：新能源与碳排放的散点矩阵", OUTPUT_DIR / "图_散点矩阵_广东.png")
        residualized_scatter(merged_gd, "emission_guangdong", "pv_gw", "广东", OUTPUT_DIR / "图_去趋势残差散点_广东_pv.png")
        residualized_scatter(merged_gd, "emission_guangdong", "wind_gw", "广东", OUTPUT_DIR / "图_去趋势残差散点_广东_wind.png")

    # Word 汇总
    doc = Document()
    doc.add_heading("新能源与碳排放关联性初步分析", level=1)
    for name in ["关联性_相关性表_九市.csv", "关联性_相关性表_广东.csv", "关联性_领先滞后相关_九市.csv", "关联性_领先滞后相关_广东.csv"]:
        p = OUTPUT_DIR / name
        if p.exists():
            df = pd.read_csv(p)
            doc.add_heading(p.stem, level=2)
            # 简化导入前10行
            head = df.head(10)
            tbl = doc.add_table(rows=1, cols=len(head.columns))
            for j, c in enumerate(head.columns):
                tbl.rows[0].cells[j].text = str(c)
            for _, r in head.iterrows():
                row = tbl.add_row().cells
                for j, c in enumerate(head.columns):
                    row[j].text = str(r[c])
    doc.save(OUTPUT_DIR / "关联性初步分析_结果汇总.docx")


# ------------------------ 主流程 ------------------------

def main():
    # 读取
    df_city = load_city_emissions()
    df_gba_year = aggregate_gba_emissions_city_to_year(df_city)
    df_gd_prov = load_apparent_emissions_guangdong()
    df_nei = load_gba_nei()

    # 中间文件
    df_city.to_csv(OUTPUT_DIR / "city_emissions_gba_1997_2019.csv", index=False)
    df_gba_year.to_csv(OUTPUT_DIR / "gba_total_emissions_1997_2019.csv", index=False)
    df_gd_prov.to_csv(OUTPUT_DIR / "guangdong_apparent_emissions_1997_2021.csv", index=False)
    df_nei.to_csv(OUTPUT_DIR / "gba_nei_2014_2023.csv", index=False)

    # 匹配窗口
    merged_gba = df_gba_year.merge(df_nei, on="year", how="inner").query("year >= 2014 and year <= 2019")
    merged_gd = df_gd_prov.merge(df_nei, on="year", how="inner").query("year >= 2014 and year <= 2021")
    panel = df_city.query("year >= 2014 and year <= 2019").copy()

    # 回归
    outputs: List[Tuple[str, sm.regression.linear_model.RegressionResultsWrapper]] = []

    # 九市时间序列
    if len(merged_gba) >= 5:
        dfg = merged_gba.copy()
        dfg["ln_emission_gba"] = safe_log(dfg["emission_gba"])
        for v in ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
            dfg[f"ln_{v}"] = safe_log(dfg[v])
        xvars = ["ln_pv_gw", "ln_wind_gw", "ln_nev_10k_units", "ln_storage_mw"]
        res_ts = fit_time_series_ols(dfg, y="ln_emission_gba", x_vars=xvars, trend=True)
        outputs.append(("九市-时间序列（2014-2019）", res_ts))

    # 广东省时间序列
    if len(merged_gd) >= 6:
        dfgd = merged_gd.copy()
        dfgd["ln_emission_gd"] = safe_log(dfgd["emission_guangdong"])
        for v in ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
            dfgd[f"ln_{v}"] = safe_log(dfgd[v])
        xvars = ["ln_pv_gw", "ln_wind_gw", "ln_nev_10k_units", "ln_storage_mw"]
        res_gd = fit_time_series_ols(dfgd, y="ln_emission_gd", x_vars=xvars, trend=True)
        outputs.append(("广东-时间序列（2014-2021）", res_gd))

    # 城市固定效应
    panel_merged = panel.merge(df_nei, on="year", how="left")
    if len(panel_merged) >= 30:
        dfp = panel_merged.copy()
        dfp["ln_emission_city"] = safe_log(dfp["emission"])
        for v in ["pv_gw", "wind_gw", "nev_10k_units", "storage_mw"]:
            dfp[f"ln_{v}"] = safe_log(dfp[v])
        dfp, city_dummies = add_city_dummies(dfp, "city_std")
        xvars = ["ln_pv_gw", "ln_wind_gw", "ln_nev_10k_units", "ln_storage_mw"] + city_dummies
        res_fe = fit_time_series_ols(dfp, y="ln_emission_city", x_vars=xvars, trend=True, cluster=dfp["year"])  # 按年份聚类稳健SE
        outputs.append(("九市-城市固定效应（2014-2019）", res_fe))

    # 文本回归结果
    with open(OUTPUT_DIR / "regression_results.txt", "w", encoding="utf-8") as f:
        for name, res in outputs:
            f.write("=" * 100 + "\n"); f.write(name + "\n"); f.write("-" * 100 + "\n")
            f.write(res.summary().as_text()); f.write("\n\n")
    # 回归系数图（主结果）
    for name, res in outputs:
        out_png = OUTPUT_DIR / f"图_回归系数_{name}.png"
        plot_coef_forest(res, name, out_png)

    # 5.1 变量选取与描述性统计
    # 九市窗口
    if not merged_gba.empty:
        stats_gba = compute_descriptive_stats(merged_gba, ["emission_gba", "pv_gw", "wind_gw", "nev_10k_units", "re_total_gw", "storage_mw"])
        stats_gba.to_csv(OUTPUT_DIR / "变量描述统计_九市_2014_2019.csv", index=False)
        save_corr_heatmap(merged_gba, ["emission_gba", "pv_gw", "wind_gw", "nev_10k_units", "storage_mw"], "九市：变量相关系数矩阵（2014-2019）", OUTPUT_DIR / "图_相关矩阵_九市_2014_2019.png")
    # 广东窗口
    if not merged_gd.empty:
        stats_gd = compute_descriptive_stats(merged_gd, ["emission_guangdong", "pv_gw", "wind_gw", "nev_10k_units", "re_total_gw", "storage_mw"])
        stats_gd.to_csv(OUTPUT_DIR / "变量描述统计_广东_2014_2021.csv", index=False)
        save_corr_heatmap(merged_gd, ["emission_guangdong", "pv_gw", "wind_gw", "nev_10k_units", "storage_mw"], "广东：变量相关系数矩阵（2014-2021）", OUTPUT_DIR / "图_相关矩阵_广东_2014_2021.png")
    # 面板窗口
    panel_merged = panel.merge(df_nei, on="year", how="left")
    if not panel_merged.empty:
        stats_panel = compute_descriptive_stats(panel_merged, ["emission", "pv_gw", "wind_gw", "nev_10k_units", "storage_mw"])
        stats_panel.to_csv(OUTPUT_DIR / "变量描述统计_面板_2014_2019.csv", index=False)

    # 5.2 弹性系数测算
    elast = elasticity_table_from_outputs(outputs)
    if not elast.empty:
        elast.to_csv(OUTPUT_DIR / "弹性系数测算.csv", index=False)

    # 5.3 稳健性检验
    robust_outs = robustness_checks(merged_gba, merged_gd, panel_merged)
    # 保存稳健性文本
    with open(OUTPUT_DIR / "robustness_results.txt", "w", encoding="utf-8") as f:
        for name, res in robust_outs:
            f.write("=" * 100 + "\n"); f.write(name + "\n"); f.write("-" * 100 + "\n")
            f.write(res.summary().as_text()); f.write("\n\n")
    # 稳健性系数图
    for name, res in robust_outs:
        out_png = OUTPUT_DIR / f"图_稳健性_回归系数_{name}.png"
        plot_coef_forest(res, name, out_png)
    # Word 汇总
    if robust_outs:
        doc = Document()
        doc.add_heading("稳健性检验结果", level=1)
        for name, res in robust_outs:
            doc.add_heading(name, level=2)
            tbl = doc.add_table(rows=1, cols=6)
            hdr = tbl.rows[0].cells
            hdr[0].text = "变量"; hdr[1].text = "系数"; hdr[2].text = "标准误"; hdr[3].text = "下界95%"; hdr[4].text = "上界95%"; hdr[5].text = "p值"
            params = res.params; se = res.bse; conf = res.conf_int(); p = res.pvalues
            for var in params.index:
                row = tbl.add_row().cells
                row[0].text = str(var)
                row[1].text = f"{params[var]:.4f}"
                row[2].text = f"{se[var]:.4f}"
                row[3].text = f"{conf.loc[var, 0]:.4f}"
                row[4].text = f"{conf.loc[var, 1]:.4f}"
                row[5].text = f"{p[var]:.4f}"
        doc.save(OUTPUT_DIR / "稳健性检验_结果汇总.docx")

    # 图表与Word表格
    plot_and_export(df_gba_year, df_nei, merged_gba, panel, outputs)

    # 4.3 两者关联性初步分析
    preliminary_association_analysis(merged_gba, merged_gd)
    # 将相关性表生成条形图
    plot_corr_table_bars(OUTPUT_DIR / "关联性_相关性表_九市.csv", "九市")
    plot_corr_table_bars(OUTPUT_DIR / "关联性_相关性表_广东.csv", "广东")

    print("已保存输出至:", OUTPUT_DIR)


if __name__ == "__main__":
    ensure_chinese_font()
    main()